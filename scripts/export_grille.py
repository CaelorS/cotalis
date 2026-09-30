#!/usr/bin/env python3
"""Exporte la grille de prix du catalogue en classeur Excel à compléter (pour le chiffreur).
Usage : python3 scripts/export_grille.py [fichier.xlsx]"""
import json, subprocess, sys, datetime, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
JS = """global.window={}; require('./assets/js/catalog.js'); const C=window.COTALIA; const items=[];
C.CATALOG.forEach(l=>l.items.forEach(it=>{items.push({lot:l.lot,id:it.id,label:it.label,sub:it.sub||'',unit:it.unit,pu:it.pu,lab:it.lab,tva:it.tva||10,nofin:!!it.nofin,qm:it.qm,qc:it.qc==null?1:it.qc,only:it.only||[],presets:Object.keys(C.PRESET).filter(k=>C.PRESET[k].includes(it.id)),why:(C.WHY&&C.WHY[it.id])||''});}));
console.log(JSON.stringify({items,modes:C.QTY_MODES,region:C.REGION,gamme:C.GAMME,notaire:C.NOTAIRE,marge:C.MARGE,fg:C.FG,pilotage:C.PILOTAGE,alea:C.ALEA,ameublement:C.AMEUBLEMENT}));"""
d = json.loads(subprocess.check_output(['node', '-e', JS], cwd=ROOT))
# valeurs modifiées dans le back-office : elles priment sur le catalogue par défaut
import urllib.request, re
cfg = open(os.path.join(ROOT, 'config.js')).read()
URL = re.search(r"supabaseUrl:\s*'([^']+)'", cfg).group(1); KEY = re.search(r"supabaseAnonKey:\s*'([^']+)'", cfg).group(1)
overrides = {}
try:
    req = urllib.request.Request(URL + '/rest/v1/pricing_items?select=*', headers={'apikey': KEY, 'Authorization': 'Bearer ' + KEY})
    overrides = {r['id']: r for r in json.loads(urllib.request.urlopen(req, timeout=15).read())}
except Exception as e:
    print('prix du back-office non lus :', e)
for it in d['items']:
    o = overrides.get(it['id'])
    if not o: it['active'] = True; continue
    it['marge'] = o.get('marge')
    for k_src, k_dst in (('pu', 'pu'), ('lab', 'lab'), ('label', 'label'), ('sub', 'sub'), ('qty_mode', 'qm'), ('qty_coef', 'qc')):
        if o.get(k_src) not in (None, ''): it[k_dst] = o[k_src]
    if o.get('tva') is not None: it['tva'] = o['tva']
    if o.get('nofin') is not None: it['nofin'] = bool(o['nofin'])
    it['active'] = o.get('active') is not False
print(len(overrides), 'ouvrage(s) modifié(s) dans le back-office repris')
try:
    req = urllib.request.Request(URL + '/rest/v1/pricing_settings?select=key,value', headers={'apikey': KEY, 'Authorization': 'Bearer ' + KEY})
    st = {r['key']: r['value'] for r in json.loads(urllib.request.urlopen(req, timeout=15).read())}
    for k_src, k_dst in (('MARGE_COUT', 'marge'), ('FG', 'fg'), ('PILOTAGE', 'pilotage'), ('ALEA', 'alea'), ('NOTAIRE', 'notaire')):
        if isinstance(st.get(k_src), (int, float)): d[k_dst] = st[k_src]
except Exception as e:
    print('réglages du back-office non lus :', e)
today = datetime.date.today()
out = sys.argv[1] if len(sys.argv) > 1 else f'cotalia-grille-prix-{today.isoformat()}.xlsx'

wb = Workbook()
F = lambda **k: Font(name='Arial', size=k.pop('size', 10), **k)
YELLOW = PatternFill('solid', start_color='FFF7CC'); HEAD = PatternFill('solid', start_color='2457A6'); LOT = PatternFill('solid', start_color='E3ECF9'); GREY = PatternFill('solid', start_color='F1F3F6')
thin = Side(style='thin', color='D3DAE1'); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
ETATS = [('bon', 'Bon'), ('correct', 'Correct'), ('degrade', 'Dégradé'), ('total', 'À rénover')]
KINDS = {'appart': 'appartement', 'maison': 'maison', 'immeuble': 'immeuble'}
modes = d['modes']

ws = wb.active; ws.title = "Mode d'emploi"
lines = [
    (f"Grille de prix Cotalia · version du {today.strftime('%d/%m/%Y')}", F(size=16, bold=True, color='2457A6')), ("", None),
    ("À quoi sert ce classeur", F(size=12, bold=True)),
    ("La feuille « Ouvrages » reprend la grille utilisée par l'estimateur en ligne. Les cases jaunes sont à compléter ou corriger ; le reste est calculé ou technique.", F()),
    ("Une fois rempli, renvoyez le fichier à Jérémy : la grille est chargée dans le back-office telle quelle, grâce à la colonne « ID technique » qu'il ne faut pas modifier.", F()), ("", None),
    ("Les colonnes de la feuille Ouvrages", F(size=12, bold=True)),
    ("Coût HT : ce que l'ouvrage nous coûte, hors taxes et avant marge, par unité indiquée (m², u, forfait…), pour une finition standard et une ville moyenne. L'estimateur applique ensuite les coefficients de zone, de finition et de complexité.", F()),
    ("Marge propre : laisser vide pour appliquer la marge par défaut de la feuille Paramètres ; saisir un pourcentage pour une marge particulière à l'ouvrage.", F()),
    ("Prix de vente HT : calculé, coût × (1 + marge). Avec la marge de 40 %, un coût de 100 € donne un prix de vente de 140 €. C'est le prix que voit le client.", F()),
    ("Part main-d'œuvre : part du prix qui est du temps de travail, en %. Le reste couvre matériaux, fournitures, location et frais. Sert aux coefficients de zone et de complexité, et à la durée du chantier.", F()),
    ("Hors finition : « oui » si le prix ne doit pas changer avec le niveau de finition (dépose, gravats, nettoyage, ouvertures de mur, câblage…).", F()),
    ("TVA : 10 % pour un logement de plus de deux ans, 5,5 % pour l'amélioration énergétique. Le prix de vente TTC se calcule tout seul.", F()),
    ("Proposé pour : états généraux du bien pour lesquels l'ouvrage est coché d'office dans l'estimation. Mettre « x » dans les colonnes voulues.", F()),
    ("Quantité proposée : règle qui calcule la quantité à partir du bien, et coefficient multiplicateur. Liste des règles dans la feuille « Paramètres ».", F()),
    ("Types de bien : laisser vide si l'ouvrage s'applique à tous, sinon lister appartement, maison, immeuble.", F()),
    ("Actif : « non » retire l'ouvrage de l'estimateur sans le supprimer.", F()), ("", None),
    ("Ajouter un ouvrage", F(size=12, bold=True)),
    ("Ajoutez une ligne à la fin, lot compris, en laissant l'ID technique vide : il sera créé au chargement. Une ligne d'exemple, grisée, montre le format attendu.", F()), ("", None),
    ("Rappel sur les marges", F(size=12, bold=True)),
    ("Depuis le 30 septembre 2026, la grille est exprimée en coûts et une marge de 40 % s'ajoute au coût par défaut. 100 % de marge double le coût.", F()),
    ("Électricité : depuis le 30 septembre 2026, trois postes distincts, tableau, appareillage, circuits et câblage. Leurs prix sont des valeurs de départ à valider.", F()),
]
for i, (t, f) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t); c.alignment = Alignment(wrap_text=True, vertical='top')
    if f: c.font = f
ws.column_dimensions['A'].width = 120

wo = wb.create_sheet('Ouvrages')
heads = ['Lot', 'Ouvrage', 'Précision', 'Unité', 'Coût HT (€)', 'Marge propre (%)', 'Prix de vente HT (€)', "Part main-d'œuvre (%)", 'Hors finition (oui/non)', 'TVA (%)', 'Prix de vente TTC (€)'] + ['Proposé : ' + e[1] for e in ETATS] + ['Quantité proposée (règle)', 'Coefficient', 'Types de bien', 'Actif (oui/non)', 'Explication affichée au client', 'ID technique']
widths = [26, 34, 30, 9, 12, 12, 14, 14, 14, 9, 14, 10, 10, 10, 10, 30, 11, 22, 11, 60, 14]
for j, h in enumerate(heads, 1):
    c = wo.cell(row=1, column=j, value=h); c.font = F(bold=True, color='FFFFFF'); c.fill = HEAD; c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center'); c.border = BORDER
    wo.column_dimensions[get_column_letter(j)].width = widths[j - 1]
wo.row_dimensions[1].height = 42
EDIT = {2, 3, 5, 6, 8, 9, 10, 12, 13, 14, 15, 16, 17, 18, 19, 20}
r = 2
def write_item(it, example=False):
    global r
    vals = [it['lot'], it['label'], it['sub'], it['unit'], it['pu'], (round(it['marge'] * 100, 1) if it.get('marge') is not None else None), f"=ROUND(E{r}*(1+IF(F{r}=\"\",'Paramètres'!$B$2,F{r}/100)),2)", round(it['lab'] * 100), 'oui' if it['nofin'] else 'non', it['tva'], f'=ROUND(G{r}*(1+J{r}/100),2)'] + ['x' if k in it['presets'] else '' for k, _ in ETATS] + [modes.get(it['qm'], it['qm']), it['qc'], ', '.join(KINDS.get(k, k) for k in it['only']), 'oui' if it.get('active', True) else 'non', it['why'], it['id']]
    for j, v in enumerate(vals, 1):
        c = wo.cell(row=r, column=j, value=v); c.font = F(color='7B8794' if example else None); c.border = BORDER
        c.alignment = Alignment(vertical='top', wrap_text=j in (2, 3, 16, 20), horizontal='center' if j in (4, 9, 10, 12, 13, 14, 15, 19) else None)
        if j in EDIT and not example: c.fill = YELLOW
        if example: c.fill = GREY
        if j in (5, 7, 11): c.number_format = '#,##0.00 €'
    r += 1
cur = None
for it in d['items']:
    if it['lot'] != cur:
        cur = it['lot']; wo.cell(row=r, column=1, value=cur).font = F(bold=True, size=11)
        for j in range(1, len(heads) + 1): wo.cell(row=r, column=j).fill = LOT
        r += 1
    write_item(it)
wo.cell(row=r, column=1, value='Exemple de ligne à ajouter (à remplacer ou supprimer)').font = F(italic=True, color='7B8794'); r += 1
write_item({'lot': 'Sols', 'label': 'Moquette', 'sub': 'pose collée, sous-couche comprise', 'unit': 'm²', 'pu': 22.86, 'marge': None, 'lab': 0.6, 'nofin': False, 'tva': 10, 'qm': 'surface', 'qc': 1, 'only': [], 'presets': ['degrade', 'total'], 'why': 'Texte court et détendu, affiché au survol de l’ouvrage dans l’estimation.', 'id': ''}, example=True)
last = r - 1
dv1 = DataValidation(type='list', formula1='"oui,non"', allow_blank=True); dv2 = DataValidation(type='list', formula1='"10,5.5,20"', allow_blank=False); dv3 = DataValidation(type='list', formula1='"' + ','.join(modes.values()).replace('"', '') + '"', allow_blank=True)
for dv in (dv1, dv2, dv3): wo.add_data_validation(dv)
dv1.add(f'I2:I{last + 40}'); dv1.add(f'S2:S{last + 40}'); dv2.add(f'J2:J{last + 40}'); dv3.add(f'P2:P{last + 40}')
wo.freeze_panes = 'C2'; wo.auto_filter.ref = f'A1:{get_column_letter(len(heads))}{last}'
wo.cell(row=1, column=6).comment = Comment("Laisser vide pour appliquer la marge par défaut (feuille Paramètres). Saisir un pourcentage pour donner à cet ouvrage une marge qui lui est propre.", 'Cotalia')
wo.cell(row=1, column=7).comment = Comment("Calculé : coût × (1 + marge). Ne pas saisir.", 'Cotalia')
wo.cell(row=1, column=11).comment = Comment("Calculé : prix de vente HT × (1 + TVA). Ne pas saisir.", 'Cotalia')
wo.cell(row=1, column=5).comment = Comment("Coût direct, avant marge, pour une finition standard en ville moyenne. Les coefficients de zone, finition et complexité s'appliquent ensuite (feuille Paramètres).", 'Cotalia')
wo.cell(row=1, column=21).comment = Comment("Identifiant utilisé par l'estimateur. Ne pas modifier. Laisser vide pour un nouvel ouvrage.", 'Cotalia')
wo.cell(row=last + 2, column=1, value='Cases jaunes : à compléter. Cases blanches : calculées ou techniques. Ligne grise : exemple.').font = F(italic=True, color='7B8794')

wp = wb.create_sheet('Paramètres')
wp.column_dimensions['A'].width = 44; wp.column_dimensions['B'].width = 14; wp.column_dimensions['C'].width = 70
rows = [('Marge', None, None, True),
    ('Marge ajoutée au coût (0,40 = +40 %)', d['marge'], 'Prix de vente = coût × (1 + marge). 1 = prix doublé.', False),
    ('Frais de notaire (% du prix d’achat, ancien)', d['notaire'], 'Utilisé dans la synthèse de financement.', False), ('', None, None, False),
    ('Coefficient de zone (appliqué à la part main-d’œuvre)', None, None, True)] + [(v[1], v[0], 'Zone « ' + k + ' »', False) for k, v in d['region'].items()] + [('', None, None, False),
    ('Coefficient de finition (appliqué à la part matériaux, sauf ouvrages « hors finition »)', None, None, True)] + [(v[1], v[0], 'Finition « ' + k + ' »', False) for k, v in d['gamme'].items()] + [('', None, None, False),
    ('Ameublement locatif (€ par m², selon stratégie)', None, None, True)] + [(k, v, 'Ajouté au coût total du projet quand la stratégie est retenue.', False) for k, v in d['ameublement'].items()] + [('', None, None, False),
    ('Règles de quantité disponibles', None, None, True)] + [(v, None, 'Code technique : ' + k, False) for k, v in modes.items()]
for i, (a, b, c_, is_head) in enumerate(rows, 1):
    if is_head:
        c = wp.cell(row=i, column=1, value=a); c.font = F(bold=True, color='FFFFFF'); c.fill = HEAD
        for j in (2, 3): wp.cell(row=i, column=j).fill = HEAD
        continue
    wp.cell(row=i, column=1, value=a).font = F()
    if b is not None:
        cb = wp.cell(row=i, column=2, value=b); cb.font = F(); cb.fill = YELLOW; cb.border = BORDER; cb.number_format = '0.00'
    if c_:
        cc = wp.cell(row=i, column=3, value=c_); cc.font = F(color='7B8794'); cc.alignment = Alignment(wrap_text=True)
wb.save(out); print('écrit :', out, '·', len(d['items']), 'ouvrages')
